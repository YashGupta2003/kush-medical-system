"""
One-time (or repeatable) bulk import of salt/composition data onto your
EXISTING master rate list. This does NOT create new medicines - it only
attaches composition text to medicines that are already in the `medicines`
table (matched by exact normalized name), so it's safe to run against your
live data.

Expected Excel format: Sheet1, row 1 = headers, from row 2 onward:
    Column A: Medicine Name   (must match a `particulars` value already in
                                the database - matched case/punctuation-
                                insensitively via the same normalize() used
                                for rate-list matching)
    Column B: Composition      (free text, e.g. "Paracetamol 650mg")

Any row whose medicine name can't be matched to an existing medicine is
written to unmatched_compositions.csv instead of being silently skipped or
guessed at with fuzzy matching - a wrong composition attached to the wrong
medicine is worse than no composition at all, so this script only matches
on an EXACT normalized name.

Usage:
    cd backend
    python scripts/backfill_composition.py /path/to/composition_list.xlsx
"""
import sys
import os
import csv

sys.path.append(os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from app.database import SessionLocal
from app.models import Medicine
from app.services.matcher import normalize
from app.services.composition_service import set_medicine_composition


def main(xlsx_path: str):
    db = SessionLocal()

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

    # Build a lookup of every existing medicine by normalized name once,
    # instead of querying the DB per row.
    by_normalized_name = {}
    for med in db.query(Medicine).all():
        by_normalized_name.setdefault(med.normalized_name, []).append(med)

    updated, skipped_blank, unmatched, ambiguous = 0, 0, [], []

    for row in ws.iter_rows(min_row=2, values_only=True):
        name_cell = row[0] if len(row) > 0 else None
        composition_cell = row[1] if len(row) > 1 else None

        if not name_cell or not str(name_cell).strip():
            skipped_blank += 1
            continue
        if not composition_cell or not str(composition_cell).strip():
            skipped_blank += 1
            continue

        name = str(name_cell).strip()
        composition = str(composition_cell).strip()
        key = normalize(name)

        candidates = by_normalized_name.get(key)
        if not candidates:
            unmatched.append((name, composition))
            continue
        if len(candidates) > 1:
            # Same name appears more than once in the master list (different
            # packs/companies) - composition is usually identical across
            # those, so apply to all of them rather than guessing which one.
            ambiguous.append((name, composition, len(candidates)))
            for med in candidates:
                set_medicine_composition(db, med.id, composition)
            updated += len(candidates)
            continue

        set_medicine_composition(db, candidates[0].id, composition)
        updated += 1

        if updated % 500 == 0:
            db.commit()
            print(f"  ...{updated} medicines updated so far")

    db.commit()
    db.close()

    print(f"Done. Updated {updated} medicines, skipped {skipped_blank} blank rows, "
          f"{len(unmatched)} unmatched, {len(ambiguous)} matched multiple rows (applied to all).")

    if unmatched:
        report_path = os.path.join(os.path.dirname(xlsx_path), "unmatched_compositions.csv")
        with open(report_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Medicine Name (not found in master list)", "Composition"])
            writer.writerows(unmatched)
        print(f"Unmatched rows written to: {report_path}")
        print("Check for spelling/punctuation differences from the master list, fix, and re-run.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python scripts/backfill_composition.py /path/to/composition_list.xlsx")
        sys.exit(1)
    main(sys.argv[1])
